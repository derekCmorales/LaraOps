from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook

from app.schemas.result import ModuleResult, NamedTable
from app.services.export_labels import label_columns, label_key, label_table_name


def module_result_to_xlsx(result: ModuleResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Solución"

    ws.append([label_key("module"), result.module])
    ws.append([label_key("status"), result.status.value])
    ws.append([])
    if result.solution.objective_value is not None:
        ws.append([label_key("objective_value"), result.solution.objective_value])
        ws.append([label_key("objective_sense"), result.solution.objective_sense])
        ws.append([])

    if result.sensitivity and result.solution.variables:
        ws.append([label_key("variable"), label_key("value"), label_key("coeff"), label_key("reduced_cost")])
        for row in _variable_rows(result):
            ws.append([row["variable"], row["value"], row["coeff"], row["reduced_cost"]])
        ws.append([])

    if result.sensitivity and result.sensitivity.constraint_analysis:
        keys = ["constraint_id", "lhs", "sense", "rhs", "slack_or_surplus", "shadow_price"]
        ws.append([label_key("constraint_analysis")])
        ws.append(label_columns(keys))
        for row in result.sensitivity.constraint_analysis:
            ws.append([row.get(k) for k in keys])
        ws.append([])
    elif result.solution.variables:
        ws.append([label_key("variable"), label_key("value")])
        for name, value in result.solution.variables.items():
            ws.append([name, value])
        ws.append([])

    ws.append([label_key("metric"), label_key("value")])
    for name, value in result.solution.metrics.items():
        ws.append([label_key(name), value])
    if result.warnings:
        ws.append([])
        ws.append([label_key("warnings")])
        for w in result.warnings:
            ws.append([w])

    if result.sensitivity is not None:
        sens_ws = wb.create_sheet("Sensibilidad")
        if result.sensitivity.constraint_analysis:
            keys = [
                "constraint_id",
                "lhs",
                "sense",
                "rhs",
                "slack_or_surplus",
                "shadow_price",
                "allowable_min_rhs",
                "allowable_max_rhs",
            ]
            sens_ws.append(label_columns(keys))
            for row in result.sensitivity.constraint_analysis:
                sens_ws.append([row.get(k) for k in keys])
            sens_ws.append([])
        if result.sensitivity.objective_ranges:
            keys = [
                "variable",
                "coeff",
                "allowable_decrease",
                "allowable_increase",
                "min_coef",
                "max_coef",
            ]
            sens_ws.append(label_columns(keys))
            for row in result.sensitivity.objective_ranges:
                sens_ws.append([row.get(k) for k in keys])
            sens_ws.append([])
        nonbasic = [
            r
            for r in result.sensitivity.reduced_costs
            if abs(float(r.get("reduced_cost", 0) or 0)) > 1e-8
        ]
        if nonbasic:
            sens_ws.append(label_columns(["variable", "reduced_cost"]))
            for row in nonbasic:
                sens_ws.append([row.get("variable"), row.get("reduced_cost")])

    if result.iterations:
        iter_ws = wb.create_sheet("Iteraciones")
        iter_ws.append(
            [label_key("index"), label_key("method"), label_key("title"), label_key("meta")]
        )
        for step in result.iterations:
            iter_ws.append([step.index, step.method, step.title, str(step.meta)])
        last = result.iterations[-1]
        if last.tableau:
            tab_ws = wb.create_sheet("Tableau final")
            for row in last.tableau:
                tab_ws.append(list(row))

    if result.tables:
        for table in result.tables:
            _write_named_table(wb, table)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _variable_rows(result: ModuleResult) -> list[dict[str, Any]]:
    sens = result.sensitivity
    if sens is None:
        return []
    rc = {str(r["variable"]): r.get("reduced_cost") for r in sens.reduced_costs}
    coef = {str(r["variable"]): r.get("coeff") for r in sens.objective_ranges}
    rows: list[dict[str, Any]] = []
    for name, value in result.solution.variables.items():
        rows.append(
            {
                "variable": name,
                "value": value,
                "coeff": coef.get(name, ""),
                "reduced_cost": rc.get(name, 0),
            }
        )
    return rows


def _write_named_table(wb: Workbook, table: NamedTable) -> None:
    title = label_table_name(table.name)[:31] or "Tabla"
    if title in wb.sheetnames:
        title = f"{title[:28]}_t"
    ws = wb.create_sheet(title)
    ws.append(label_columns(list(table.columns)))
    for row in table.rows:
        ws.append(list(row))
