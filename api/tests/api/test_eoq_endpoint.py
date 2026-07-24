from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import assert_allclose

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_eoq_solve_endpoint(client):
    response = client.post(
        "/api/v1/modules/eoq/solve",
        json={"D": 1000, "S": 10, "H": 0.5},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["module"] == "eoq"
    assert body["status"] == "ok"
    assert body["iterations"] is None
    assert body["sensitivity"] is None
    assert_allclose(body["solution"]["variables"]["Q"], 200.0)
    assert_allclose(body["solution"]["metrics"]["Q_star"], 200.0)
    assert_allclose(body["solution"]["metrics"]["TC"], 100.0)
    assert_allclose(body["solution"]["metrics"]["orders_per_year"], 5.0)
    assert_allclose(body["solution"]["metrics"]["TC_ordering"], 50.0)
    assert_allclose(body["solution"]["metrics"]["TC_holding"], 50.0)
    series_names = {s["name"] for s in body["graph"]["series"]}
    assert series_names == {"TC", "ordering", "holding"}


def test_eoq_export_xlsx(client):
    response = client.post(
        "/api/v1/modules/eoq/export.xlsx",
        json={"D": 1000, "S": 10, "H": 0.5},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA)
    wb = load_workbook(BytesIO(response.content))
    assert "Solución" in wb.sheetnames


def test_eoq_export_pdf(client):
    from pypdf import PdfReader

    response = client.post(
        "/api/v1/modules/eoq/export.pdf",
        json={"D": 1000, "S": 10, "H": 0.5},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    reader = PdfReader(BytesIO(response.content))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "LaraOps" in text
    assert "eoq" in text.lower()


def test_eoq_invalid_d_zero_422(client):
    response = client.post(
        "/api/v1/modules/eoq/solve",
        json={"D": 0, "S": 10, "H": 0.5},
    )
    assert response.status_code == 422


def test_eoq_invalid_negative_s_422(client):
    response = client.post(
        "/api/v1/modules/eoq/solve",
        json={"D": 1000, "S": -1, "H": 0.5},
    )
    assert response.status_code == 422


def test_eoq_missing_h_422(client):
    response = client.post(
        "/api/v1/modules/eoq/solve",
        json={"D": 1000, "S": 10},
    )
    assert response.status_code == 422
